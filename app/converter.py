from collections.abc import Callable
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt

from app.config import MAX_PDF_PAGES, OCR_LANGUAGES
from app.errors import ConversionError
from app.extraction import PageContent, extract_page


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    column_count = max((len(row) for row in rows), default=0)
    if column_count == 0:
        return
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value


def _add_page_content(document: Document, content: PageContent) -> None:
    for item in content.items:
        if item.kind == "paragraph":
            paragraph = document.add_paragraph()
            for line_index, line in enumerate(item.value.splitlines()):
                if line_index:
                    paragraph.add_run().add_break()
                paragraph.add_run(line)
        elif item.kind == "table":
            _add_table(document, item.value)
        elif item.kind == "image":
            try:
                document.add_picture(BytesIO(item.value), width=Cm(15.5))
            except Exception:
                continue


def convert_pdf(
    input_path: Path,
    output_path: Path,
    progress: Callable[[int, int], None],
    max_pages: int = MAX_PDF_PAGES,
) -> None:
    input_path = Path(input_path)
    output_path = Path(output_path)
    temporary_output = output_path.with_suffix(".docx.part")
    temporary_output.unlink(missing_ok=True)

    try:
        source = pymupdf.open(input_path)
    except Exception as exc:
        raise ConversionError("PDF 文件已损坏或无法打开") from exc

    try:
        if source.needs_pass:
            raise ConversionError("PDF 已加密，请先移除密码后再转换")
        total_pages = source.page_count
        if total_pages == 0:
            raise ConversionError("PDF 没有可转换的页面")
        if total_pages > max_pages:
            raise ConversionError(f"PDF 页数过多，最多支持 {max_pages} 页")

        word = Document()
        _configure_document(word)
        for page_index, page in enumerate(source):
            content = extract_page(page, OCR_LANGUAGES)
            _add_page_content(word, content)
            if page_index < total_pages - 1:
                word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            progress(page_index + 1, total_pages)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        word.save(temporary_output)
        temporary_output.replace(output_path)
    except ConversionError:
        temporary_output.unlink(missing_ok=True)
        raise
    except Exception as exc:
        temporary_output.unlink(missing_ok=True)
        raise ConversionError("转换失败，PDF 页面可能包含不受支持的内容") from exc
    finally:
        source.close()
