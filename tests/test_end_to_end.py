import shutil
import subprocess

import pymupdf
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from app.converter import convert_pdf


CHINESE_FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"


def _has_ocr_languages() -> bool:
    if shutil.which("tesseract") is None:
        return False
    result = subprocess.run(
        ["tesseract", "--list-langs"],
        capture_output=True,
        text=True,
        check=False,
    )
    languages = set(result.stdout.splitlines())
    return {"chi_sim", "eng"}.issubset(languages)


@pytest.mark.skipif(
    not _has_ocr_languages(),
    reason="需要 Tesseract 的 chi_sim 和 eng 语言包",
)
def test_mixed_pdf_converts_text_and_scanned_pages(tmp_path):
    image_path = tmp_path / "scan.png"
    image = Image.new("RGB", (1600, 1000), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(CHINESE_FONT, 104)
    draw.text((130, 320), "扫描文档测试", fill="black", font=font)
    image.save(image_path, dpi=(200, 200))

    pdf_path = tmp_path / "mixed.pdf"
    source = pymupdf.open()
    text_page = source.new_page()
    text_page.insert_text((72, 72), "Normal editable page")
    scan_page = source.new_page()
    scan_page.insert_image(scan_page.rect, filename=str(image_path))
    source.save(pdf_path)
    source.close()

    output_path = tmp_path / "mixed.docx"
    progress = []
    convert_pdf(
        pdf_path,
        output_path,
        lambda current, total: progress.append((current, total)),
    )

    word = Document(output_path)
    text = "".join(paragraph.text.replace(" ", "") for paragraph in word.paragraphs)
    assert "Normaleditablepage" in text
    assert "文档" in text
    assert progress == [(1, 2), (2, 2)]


@pytest.mark.skipif(
    not _has_ocr_languages(),
    reason="需要 Tesseract 的 chi_sim 和 eng 语言包",
)
def test_english_scanned_pdf_becomes_editable_text(tmp_path):
    image_path = tmp_path / "english-scan.png"
    image = Image.new("RGB", (1600, 1000), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        104,
    )
    draw.text((110, 320), "English document", fill="black", font=font)
    image.save(image_path, dpi=(200, 200))

    pdf_path = tmp_path / "english-scan.pdf"
    source = pymupdf.open()
    page = source.new_page()
    page.insert_image(page.rect, filename=str(image_path))
    source.save(pdf_path)
    source.close()

    output_path = tmp_path / "english-scan.docx"
    convert_pdf(pdf_path, output_path, lambda *_: None)
    word = Document(output_path)
    text = " ".join(paragraph.text for paragraph in word.paragraphs)
    assert "English document" in text
