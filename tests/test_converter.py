import pymupdf
import pytest
from docx import Document

from app.converter import convert_pdf
from app.errors import ConversionError


def test_text_pdf_becomes_editable_docx(tmp_path):
    pdf_path = tmp_path / "input.pdf"
    source = pymupdf.open()
    first = source.new_page()
    first.insert_text((72, 72), "Editable PDF text")
    second = source.new_page()
    second.insert_text((72, 72), "Second page text")
    source.save(pdf_path)
    source.close()

    output = tmp_path / "output.docx"
    calls = []
    convert_pdf(pdf_path, output, lambda current, total: calls.append((current, total)))

    word = Document(output)
    all_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "Editable PDF text" in all_text
    assert "Second page text" in all_text
    assert calls == [(1, 2), (2, 2)]


def test_empty_pdf_is_rejected_without_partial_output(tmp_path):
    pdf_path = tmp_path / "empty.pdf"
    objects = [
        b"1 0 obj\n<</Type /Catalog /Pages 2 0 R>>\nendobj\n",
        b"2 0 obj\n<</Type /Pages /Count 0 /Kids []>>\nendobj\n",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for item in objects:
        offsets.append(len(pdf))
        pdf.extend(item)
    xref_offset = len(pdf)
    pdf.extend(b"xref\n0 3\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        b"trailer\n<</Size 3 /Root 1 0 R>>\nstartxref\n"
        + str(xref_offset).encode()
        + b"\n%%EOF\n"
    )
    pdf_path.write_bytes(pdf)
    output = tmp_path / "output.docx"

    with pytest.raises(ConversionError, match="没有可转换的页面"):
        convert_pdf(pdf_path, output, lambda *_: None)

    assert not output.exists()


def test_pdf_over_page_limit_is_rejected_before_conversion(tmp_path):
    pdf_path = tmp_path / "too-many-pages.pdf"
    source = pymupdf.open()
    source.new_page()
    source.new_page()
    source.save(pdf_path)
    source.close()
    output = tmp_path / "output.docx"

    with pytest.raises(ConversionError, match="最多支持 1 页"):
        convert_pdf(pdf_path, output, lambda *_: None, max_pages=1)

    assert not output.exists()


def test_pdf_table_becomes_one_editable_word_table_without_duplicate_paragraphs(tmp_path):
    pdf_path = tmp_path / "table.pdf"
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text((72, 50), "Report heading")
    x_positions = [72, 200, 330]
    y_positions = [90, 130, 170]
    for x in x_positions:
        page.draw_line((x, y_positions[0]), (x, y_positions[-1]))
    for y in y_positions:
        page.draw_line((x_positions[0], y), (x_positions[-1], y))
    page.insert_text((82, 116), "Name")
    page.insert_text((210, 116), "Value")
    page.insert_text((82, 156), "Alpha")
    page.insert_text((210, 156), "42")
    source.save(pdf_path)
    source.close()

    output = tmp_path / "table.docx"
    convert_pdf(pdf_path, output, lambda *_: None)
    word = Document(output)

    assert len(word.tables) == 1
    assert word.tables[0].cell(0, 0).text == "Name"
    assert word.tables[0].cell(1, 1).text == "42"
    paragraph_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "Report heading" in paragraph_text
    assert "Alpha" not in paragraph_text
